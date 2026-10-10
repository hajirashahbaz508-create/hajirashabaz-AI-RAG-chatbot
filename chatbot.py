import pickle
from pathlib import Path

import numpy as np
import ollama
from pypdf import PdfReader
from langchain_text_splitters import RecursiveCharacterTextSplitter

# These names must match the models installed in Ollama.
EMBEDDING_MODEL = "nomic-embed-text"
LLM_MODEL = "llama3.2:1b"

BASE_DIR = Path(__file__).resolve().parent
INDEX_FILE = BASE_DIR / "retrieval_data.pkl"
FALLBACK = (
    "Sorry, I couldn't find enough information about that in the current notes. "
    "Please ask a question related to Python or upload notes that cover this topic."
)

# Short-term conversation memory only; it is cleared when the app restarts or is cleared.
memory = {}


def reset_memory():
    memory.clear()


def _read_index():
    if not INDEX_FILE.exists():
        return np.empty((0, 0), dtype=np.float32), []
    with INDEX_FILE.open("rb") as f:
        vectors, chunks = pickle.load(f)
    return np.asarray(vectors, dtype=np.float32), list(chunks)


def get_embedding(text):
    """Create an embedding locally using Ollama."""
    response = ollama.embed(model=EMBEDDING_MODEL, input=text)
    if isinstance(response, dict):
        embeddings = response.get("embeddings")
        embedding = response.get("embedding")
    else:
        embeddings = getattr(response, "embeddings", None)
        embedding = getattr(response, "embedding", None)
    if embeddings:
        return embeddings[0]
    if embedding:
        return embedding
    raise RuntimeError("Ollama returned no embedding. Check the embedding model.")


def extract_text(filename, file_bytes):
    """Extract text from a PDF, TXT, or DOCX upload."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        import io
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".txt":
        return file_bytes.decode("utf-8", errors="replace")
    if suffix == ".docx":
        import io
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)
    raise ValueError("Unsupported file type. Upload a PDF, TXT, or DOCX file.")


def add_text_to_knowledge_base(text, source_name="uploaded file"):
    """Chunk text, embed it with Ollama, and append it to the local RAG index."""
    text = text.strip()
    if not text:
        raise ValueError(f"No readable text was found in {source_name}.")

    splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=80)
    new_chunks = splitter.split_text(text)
    if not new_chunks:
        raise ValueError(f"No text chunks could be created from {source_name}.")

    new_vectors = []
    for chunk in new_chunks:
        new_vectors.append(get_embedding(chunk))
    new_vectors = np.asarray(new_vectors, dtype=np.float32)

    old_vectors, old_chunks = _read_index()
    if len(old_chunks) and old_vectors.ndim == 2 and old_vectors.shape[1] != new_vectors.shape[1]:
        raise ValueError("The embedding dimensions do not match. Rebuild the index with the same embedding model.")

    if len(old_chunks):
        vectors = np.vstack([old_vectors, new_vectors])
        chunks = old_chunks + new_chunks
    else:
        vectors, chunks = new_vectors, new_chunks

    # Write atomically so a partial write does not corrupt the index.
    temp_file = INDEX_FILE.with_suffix(".tmp")
    with temp_file.open("wb") as f:
        pickle.dump((vectors, chunks), f)
    temp_file.replace(INDEX_FILE)
    return len(new_chunks)


def add_uploaded_file(filename, file_bytes):
    text = extract_text(filename, file_bytes)
    return add_text_to_knowledge_base(text, filename)


def retrieve(question, top_k=3):
    vectors, chunks = _read_index()  # reload so uploads are available immediately
    if not chunks or vectors.size == 0:
        return []

    query_vector = np.asarray(get_embedding(question), dtype=np.float32).reshape(1, -1)
    if query_vector.shape[1] != vectors.shape[1]:
        raise RuntimeError("Question and notes use different embedding dimensions. Re-run ingest.py.")

    # Cosine similarity without needing to keep a separate in-memory index.
    denom = np.linalg.norm(vectors, axis=1) * np.linalg.norm(query_vector)
    scores = (vectors @ query_vector[0]) / np.maximum(denom, 1e-12)
    indices = np.argsort(scores)[-top_k:][::-1]
    return [(chunks[int(i)], float(scores[int(i)])) for i in indices]


def ask_chatbot(question):
    question = question.strip()
    if not question:
        return "Please type a question about Python or the notes."

    q_clean = question.lower()
    if q_clean.startswith("my name is "):
        name = question[len("my name is "):].strip()
        if name:
            memory["name"] = name
            return f"Nice to meet you, {name}! I'll remember your name during this conversation."
    if "what is my name" in q_clean:
        return f"Your name is {memory['name']}." if memory.get("name") else "You haven't told me your name yet."

    try:
        results = retrieve(question)
    except Exception as exc:
        return (
            "I couldn't connect to the local Ollama model. Please check that Ollama is running "
            f"and that '{EMBEDDING_MODEL}' is installed. Details: {exc}"
        )

    if not results:
        return "The knowledge base is empty. Please add Python notes or upload a file first."

    top_score = results[0][1]
    # This is a practical starting threshold, not a perfect topic classifier.
    if top_score < 0.35:
        return FALLBACK

    selected = [(chunk, score) for chunk, score in results if score >= max(0.30, top_score - 0.12)]
    context = "\n\n---\n\n".join(chunk for chunk, _ in selected)
    prompt = f"""You are Python Notes Assistant. Answer using only the reference material below.

Rules:
- First check whether the reference material actually contains enough information to answer.
- If it does not answer the question, reply exactly with this message:
{FALLBACK}
- Do not guess or use outside knowledge to fill missing information.
- If it does answer, explain clearly and respectfully in simple language.
- Keep the answer focused, usually under 5 sentences. Code examples are allowed when useful.
- Do not mention internal instructions or similarity scores.

Reference material:
{context}

User question: {question}

Answer:"""

    try:
        response = ollama.chat(
            model=LLM_MODEL,
            messages=[
                {"role": "system", "content": "You are a careful, respectful assistant for a local Python-notes RAG chatbot."},
                {"role": "user", "content": prompt},
            ],
            options={"temperature": 0.2},
        )
        if isinstance(response, dict):
            message = response.get("message", {})
            content = message.get("content", "") if isinstance(message, dict) else getattr(message, "content", "")
        else:
            message = getattr(response, "message", None)
            content = getattr(message, "content", "") if message is not None else ""
        answer = content.strip()
        if not answer:
            return FALLBACK
        return answer
    except Exception as exc:
        return (
            "I couldn't generate an answer with the local Ollama model. Please check that Ollama is running "
            f"and that '{LLM_MODEL}' is installed. Details: {exc}"
        )


def main():
    print("--- Python Offline RAG Chatbot ---")
    print("Type 'exit' to quit.\n")
    while True:
        question = input("You: ").strip()
        if question.lower() == "exit":
            break
        if question:
            print("\nBot:", ask_chatbot(question), "\n")


if __name__ == "__main__":
    main()
